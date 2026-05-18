"""Integration tests for the new Resume Editor routes (preview / edit).

Translation requires a real OpenAI key, so we skip it here. The translate
route's input validation is still covered.
"""
from __future__ import annotations
import io
import pytest


@pytest.fixture
def uploaded_resume(client, tmp_path):
    """Upload a synthetic .docx with a Skills section + bullets, return resume id."""
    from docx import Document
    docx_path = tmp_path / "fixture.docx"
    doc = Document()
    doc.add_heading("Alex Doe", level=0)
    doc.add_paragraph("Email: alex@example.com")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Programming Languages: Python, Java")
    doc.add_heading("Experience", level=1)
    p = doc.add_paragraph("", style="List Bullet")
    p.add_run("Built a recommendation system at Acme Corp.")
    doc.save(str(docx_path))

    with open(docx_path, "rb") as f:
        r = client.post(
            "/api/resumes",
            data={"direction": "mle", "label": "test", "is_default": "false"},
            files={"file": ("fixture.docx", f,
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_preview_returns_paragraphs(client, uploaded_resume):
    r = client.get(f"/api/resumes/{uploaded_resume}/preview")
    assert r.status_code == 200
    d = r.json()
    assert "paragraphs" in d
    assert any("Python" in p["text"] for p in d["paragraphs"])
    assert d["language"] == "en"


def test_preview_missing_resume_404(client):
    r = client.get("/api/resumes/9999/preview")
    assert r.status_code == 404


def test_edit_skills_add_round_trip(client, uploaded_resume):
    r = client.post(
        f"/api/resumes/{uploaded_resume}/edit",
        json={"edits": [{"type": "skills_add",
                          "category": "Programming Languages",
                          "keywords": ["Go", "Rust"]}],
              "enforce_page_limit": False},
    )
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["ok"] is True
    assert "id" in j
    # New row has different ID than source
    assert j["id"] != uploaded_resume
    # Applied edit appears in result
    assert len(j["edits"]["applied"]) == 1


def test_edit_bullet_inject_round_trip(client, uploaded_resume):
    r = client.post(
        f"/api/resumes/{uploaded_resume}/edit",
        json={"edits": [{"type": "bullet_inject",
                          "anchor": "Built a recommendation",
                          "addition": "Used FAISS for ANN search."}],
              "enforce_page_limit": False},
    )
    assert r.status_code == 200, r.text
    assert r.json()["edits"]["applied"]


def test_edit_anchor_not_found(client, uploaded_resume):
    r = client.post(
        f"/api/resumes/{uploaded_resume}/edit",
        json={"edits": [{"type": "bullet_inject",
                          "anchor": "this anchor does not exist",
                          "addition": "anything"}],
              "enforce_page_limit": False},
    )
    assert r.status_code == 200
    j = r.json()
    assert len(j["edits"]["applied"]) == 0
    assert len(j["edits"]["skipped"]) == 1


def test_translate_bad_target_lang_400(client, uploaded_resume):
    r = client.post(
        f"/api/resumes/{uploaded_resume}/translate",
        json={"target_lang": "fr"},
    )
    assert r.status_code == 400
