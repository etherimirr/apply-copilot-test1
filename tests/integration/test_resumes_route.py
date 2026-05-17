"""Integration: resume upload / list / download / delete via TestClient."""
from __future__ import annotations
import io


def _make_fake_docx():
    """Build a tiny in-memory .docx (just minimum bytes — we're not parsing it,
    just storing/serving it back). python-docx-built valid file."""
    from docx import Document
    d = Document()
    d.add_heading("Test Resume", level=1)
    d.add_paragraph("Lorem ipsum.")
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf.read()


def test_list_initially_empty(client):
    r = client.get("/api/resumes")
    assert r.status_code == 200
    assert r.json() == []


def test_upload_lists_downloads_deletes(client, tmp_path):
    content = _make_fake_docx()
    r = client.post(
        "/api/resumes",
        data={"direction": "mle", "label": "ML v1", "is_default": "true"},
        files={"file": ("my_resume.docx", content,
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200, r.text
    rid = r.json()["id"]

    # List
    items = client.get("/api/resumes").json()
    assert len(items) == 1
    assert items[0]["direction"] == "mle"
    assert items[0]["is_default"] is True

    # Download
    r2 = client.get(f"/api/resumes/{rid}/download")
    assert r2.status_code == 200
    assert r2.content == content   # round-trip preserves bytes

    # Delete
    r3 = client.delete(f"/api/resumes/{rid}")
    assert r3.status_code == 200
    assert client.get("/api/resumes").json() == []


def test_upload_rejects_non_docx(client):
    r = client.post(
        "/api/resumes",
        data={"direction": "mle", "label": "x"},
        files={"file": ("foo.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert r.status_code == 400


def test_upload_rejects_unknown_direction(client):
    r = client.post(
        "/api/resumes",
        data={"direction": "marketing", "label": "x"},
        files={"file": ("r.docx", _make_fake_docx(), "application/octet-stream")},
    )
    assert r.status_code == 400


def test_uploading_default_unsets_prior_default(client):
    content = _make_fake_docx()
    client.post("/api/resumes",
                 data={"direction": "mle", "is_default": "true"},
                 files={"file": ("a.docx", content, "application/octet-stream")})
    client.post("/api/resumes",
                 data={"direction": "sde_general", "is_default": "true"},
                 files={"file": ("b.docx", content, "application/octet-stream")})

    items = client.get("/api/resumes").json()
    defaults = [r for r in items if r["is_default"]]
    assert len(defaults) == 1
    assert defaults[0]["direction"] == "sde_general"
