"""End-to-end resume editor tests.

Crafts a fixture .docx programmatically, applies edits, verifies:
- Bold-only-on-added rule (existing text stays unbold)
- Anchor finding (skills + bullet)
- Page-limit rollback path (when a giant edit blows past max_pages)

Skips the PDF render path if no PDF converter is available (CI environments
without LibreOffice / Word / docx2pdf-supporting-stack).
"""
from __future__ import annotations
import importlib
import shutil
from pathlib import Path

import pytest


def _has_pdf_converter() -> bool:
    """True if any PDF conversion path is likely to work."""
    if shutil.which("soffice") or shutil.which("libreoffice"):
        return True
    if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists():
        return True
    try:
        import docx2pdf  # noqa: F401
        return True
    except ImportError:
        return False


@pytest.fixture
def fixture_resume(tmp_path) -> Path:
    """A minimal .docx with a Skills section and an Experience bullet."""
    from docx import Document
    doc = Document()
    doc.add_heading("Jane Doe", level=0)
    doc.add_paragraph("Email: jane@example.com | LinkedIn: linkedin.com/in/jane")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Programming Languages: Python, Java")
    doc.add_paragraph("ML Frameworks: PyTorch")

    doc.add_heading("Experience", level=1)
    p = doc.add_paragraph("", style="List Bullet")
    p.add_run("Built a recommendation system handling 1M users at Acme Corp.")
    p2 = doc.add_paragraph("", style="List Bullet")
    p2.add_run("Shipped ML pipelines to production with FastAPI and Docker.")

    out = tmp_path / "fixture.docx"
    doc.save(str(out))
    return out


def test_load_and_text_extract(fixture_resume):
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    text = resume_editor.docx_to_text(doc)
    assert "Python, Java" in text
    assert "PyTorch" in text
    assert "Acme Corp" in text


def test_skills_add_appends_and_bolds_only_new(fixture_resume):
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    ok = resume_editor.skills_add(doc, "Programming Languages", ["Go", "Rust"])
    assert ok is True

    # Find the Programming Languages paragraph
    pl_para = next(p for p in doc.paragraphs
                    if p.text.startswith("Programming Languages:"))
    # Original text and new keywords both appear
    assert "Python" in pl_para.text
    assert "Go" in pl_para.text
    assert "Rust" in pl_para.text

    # Bold-only-on-added: original runs should not be bold; added runs should be.
    bold_texts = [r.text for r in pl_para.runs if r.bold]
    nonbold_texts = [r.text for r in pl_para.runs if not r.bold]
    assert any("Go" == t for t in bold_texts)
    assert any("Rust" == t for t in bold_texts)
    # Python should remain in a non-bold run (it was part of the original)
    assert any("Python" in t for t in nonbold_texts)


def test_skills_add_no_match_returns_false(fixture_resume):
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    ok = resume_editor.skills_add(doc, "Cloud Infrastructure", ["AWS"])
    # No "Cloud Infrastructure:" paragraph exists → returns False
    assert ok is False


def test_bullet_inject_finds_anchor(fixture_resume):
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    ok = resume_editor.bullet_inject(doc, "Built a recommendation",
                                          "Used FAISS for ANN search.")
    assert ok is True

    bullet = next(p for p in doc.paragraphs
                   if p.text.startswith("Built a recommendation"))
    assert "FAISS" in bullet.text
    # Newly-added run should be bold
    bold_texts = [r.text for r in bullet.runs if r.bold]
    assert any("FAISS" in t for t in bold_texts)


def test_bullet_inject_anchor_missing_returns_false(fixture_resume):
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    ok = resume_editor.bullet_inject(doc, "Something that doesn't exist",
                                          "Anything.")
    assert ok is False


def test_apply_edits_with_rollback_applies_all_when_short(fixture_resume, tmp_path):
    """When the resume stays under the page limit, all edits should apply."""
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    edits = [
        {"type": "skills_add", "category": "Programming Languages",
         "keywords": ["Go"]},
        {"type": "bullet_inject", "anchor": "Built a recommendation",
         "addition": "Used FAISS for ANN search."},
    ]
    out_pdf = tmp_path / "out.pdf"
    result = resume_editor.apply_edits_with_rollback(doc, edits, out_pdf,
                                                          max_pages=10)
    assert len(result["applied"]) == 2
    assert len(result["skipped"]) == 0


@pytest.mark.skipif(not _has_pdf_converter(),
                     reason="No PDF converter (LibreOffice/Word) available")
def test_apply_edits_rollback_triggers_when_over_page_limit(fixture_resume, tmp_path):
    """If max_pages=0 the rollback path always fires.

    Tests that the doc's existing content is preserved (rollback works) and
    nothing leaks past."""
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    before_text = resume_editor.docx_to_text(doc)
    edits = [
        {"type": "skills_add", "category": "Programming Languages",
         "keywords": ["Go", "Rust", "Kotlin"]},
    ]
    out_pdf = tmp_path / "out_rollback.pdf"
    # max_pages=0 forces every edit to "exceed" → rollback
    result = resume_editor.apply_edits_with_rollback(doc, edits, out_pdf,
                                                          max_pages=0)
    assert len(result["applied"]) == 0
    assert len(result["skipped"]) == 1
    # Doc was reverted: text is unchanged
    after_text = resume_editor.docx_to_text(doc)
    assert "Go" not in after_text
    assert before_text in after_text or after_text == before_text


def test_save_and_reload_preserves_edits(fixture_resume, tmp_path):
    """Edit, save, reload — edits persist (basic docx round-trip sanity)."""
    from apply_copilot.core import resume_editor
    doc = resume_editor.load_docx(fixture_resume)
    resume_editor.skills_add(doc, "Programming Languages", ["Go"])
    out = tmp_path / "edited.docx"
    resume_editor.save_docx(doc, out)

    doc2 = resume_editor.load_docx(out)
    text2 = resume_editor.docx_to_text(doc2)
    assert "Go" in text2
    assert "Python" in text2  # original kept
